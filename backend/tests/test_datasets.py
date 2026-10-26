import hashlib
import io
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from zipfile import ZipFile

import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq
from django.conf import settings
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import IntegrityError, transaction
from django.test import Client, TestCase, override_settings
from django.utils.dateparse import parse_datetime

from datasets.models import Dataset


@override_settings(ALLOWED_HOSTS=["testserver"])
class DatasetApiTests(TestCase):
    endpoint = "/api/datasets/"
    public_fields = {"id", "filename", "format", "size_bytes", "uploaded_at"}

    def setUp(self):
        media = TemporaryDirectory(prefix="cloudetl-dataset-test-")
        self.addCleanup(media.cleanup)
        self.media_root = Path(media.name)
        media_settings = override_settings(MEDIA_ROOT=self.media_root)
        media_settings.enable()
        self.addCleanup(media_settings.disable)
        user_model = get_user_model()
        self.user = user_model.objects.create_user(username="dataset-owner")
        self.other_user = user_model.objects.create_user(username="other-owner")
        self.client = Client(enforce_csrf_checks=True)
        self.client.force_login(self.user)
        self.token = self.csrf_token(self.client)

    def csrf_token(self, client):
        response = client.get("/api/auth/csrf/")
        self.assertEqual(response.status_code, 200)
        self.assertIn("csrftoken", client.cookies)
        return response.json()["csrfToken"]

    def post_file(self, filename="source.csv", content=b"name,value\nriver,1\n", **extra):
        upload = SimpleUploadedFile(filename, content, content_type="application/octet-stream")
        return self.client.post(
            self.endpoint, {"file": upload, **extra},
            HTTP_X_CSRFTOKEN=self.token, HTTP_ORIGIN="http://testserver",
        )

    def stored_files(self):
        return {path.relative_to(self.media_root) for path in self.media_root.rglob("*") if path.is_file()}

    def assert_public_metadata(self, payload, dataset):
        self.assertEqual(set(payload), self.public_fields)
        self.assertEqual(payload["id"], str(dataset.pk))
        self.assertEqual(payload["filename"], dataset.filename)
        self.assertEqual(payload["format"], dataset.format)
        self.assertEqual(payload["size_bytes"], dataset.size_bytes)
        self.assertEqual(parse_datetime(payload["uploaded_at"]), dataset.uploaded_at)

    def assert_invalid_upload(self, filename, content):
        count = Dataset.objects.count()
        stored_before = self.stored_files()
        response = self.post_file(filename, content)
        self.assertEqual(response.status_code, 400, response.content)
        error = response.json()["error"]
        self.assertEqual(error["code"], "validation_error")
        self.assertTrue(error["fields"]["file"])
        self.assertEqual(Dataset.objects.count(), count)
        self.assertEqual(self.stored_files(), stored_before)
        return response

    def valid_formats(self):
        workbook_buffer = io.BytesIO()
        workbook = openpyxl.Workbook()
        workbook.active.append(["name", "value"])
        workbook.active.append(["river", 1])
        workbook.save(workbook_buffer)
        workbook.close()

        parquet_buffer = io.BytesIO()
        pq.write_table(pa.table({"name": ["river"], "value": [1]}), parquet_buffer)
        return {
            "csv": b"\xef\xbb\xbfname,value\r\nriver,1\r\n",
            "tsv": b"name\tvalue\nriver\t1\n",
            "json": b'[{"name":"river","value":1}]\n',
            "xlsx": workbook_buffer.getvalue(),
            "parquet": parquet_buffer.getvalue(),
        }

    def test_all_supported_formats_preserve_original_bytes_and_private_hash(self):
        for extension, content in self.valid_formats().items():
            with self.subTest(extension=extension):
                response = self.post_file(f"source.{extension}", content)
                self.assertEqual(response.status_code, 201, response.content)
                self.assertEqual(set(response.json()), {"dataset"})
                payload = response.json()["dataset"]
                dataset = Dataset.objects.get(pk=payload["id"])
                self.assertEqual(dataset.owner, self.user)
                self.assertEqual(dataset.filename, f"source.{extension}")
                self.assertEqual(dataset.format, extension)
                self.assertEqual(dataset.size_bytes, len(content))
                self.assertEqual(dataset.sha256, hashlib.sha256(content).hexdigest())
                with dataset.file.open("rb") as stored:
                    self.assertEqual(stored.read(), content)
                self.assertEqual(dataset.file.name, f"datasets/{self.user.pk}/{dataset.pk}.{extension}")
                self.assert_public_metadata(payload, dataset)
                self.assertNotIn(dataset.file.name, response.content.decode())
                self.assertNotIn(dataset.sha256, response.content.decode())
                self.assertIn("no-store", response.headers["Cache-Control"])

        response = self.client.get(self.endpoint)
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(set(data), {"datasets", "page", "total_pages", "total_count", "limits"})
        self.assertEqual(data["total_count"], 5)
        self.assertEqual(data["limits"], {
            "max_file_size_bytes": settings.DATASET_MAX_FILE_SIZE_BYTES,
            "extensions": ["csv", "tsv", "json", "xlsx", "parquet"],
        })
        for payload in data["datasets"]:
            self.assert_public_metadata(payload, Dataset.objects.get(pk=payload["id"]))
        self.assertIn("no-store", response.headers["Cache-Control"])

    def test_list_is_empty_only_for_current_owner_and_never_leaks_other_users(self):
        created = self.post_file("private.csv")
        self.assertEqual(created.status_code, 201)
        own_id = created.json()["dataset"]["id"]
        other_client = Client(enforce_csrf_checks=True)
        other_client.force_login(self.other_user)
        other_token = self.csrf_token(other_client)
        empty = other_client.get(self.endpoint)
        self.assertEqual(empty.status_code, 200)
        self.assertEqual(empty.json()["datasets"], [])
        self.assertEqual(empty.json()["total_count"], 0)
        self.assertEqual(empty.json()["total_pages"], 1)
        other = other_client.post(
            self.endpoint, {"file": SimpleUploadedFile("other.csv", b"id\n2\n")},
            HTTP_X_CSRFTOKEN=other_token, HTTP_ORIGIN="http://testserver",
        )
        self.assertEqual(other.status_code, 201)
        own_list = self.client.get(self.endpoint).json()
        other_list = other_client.get(self.endpoint).json()
        self.assertEqual([item["id"] for item in own_list["datasets"]], [own_id])
        self.assertEqual([item["filename"] for item in other_list["datasets"]], ["other.csv"])
        self.assertEqual(own_list["total_count"], 1)
        self.assertEqual(other_list["total_count"], 1)
        self.assertNotIn(own_id, str(other_list))

    def test_anonymous_list_and_upload_require_session_even_with_valid_csrf(self):
        anonymous = Client(enforce_csrf_checks=True)
        token = self.csrf_token(anonymous)
        response = anonymous.get(self.endpoint)
        self.assertEqual(response.status_code, 401)
        self.assertEqual(response.json()["error"]["code"], "authentication_required")
        response = anonymous.post(
            self.endpoint, {"file": SimpleUploadedFile("source.csv", b"id\n1\n")},
            HTTP_X_CSRFTOKEN=token, HTTP_ORIGIN="http://testserver",
        )
        self.assertEqual(response.status_code, 401)
        self.assertEqual(response.json()["error"]["code"], "authentication_required")
        self.assertEqual(Dataset.objects.count(), 0)
        self.assertEqual(self.stored_files(), set())

    def test_upload_requires_csrf_token_and_trusted_origin(self):
        for headers in (
            {},
            {"HTTP_X_CSRFTOKEN": self.token, "HTTP_ORIGIN": "https://attacker.example"},
        ):
            with self.subTest(headers=sorted(headers)):
                response = self.client.post(
                    self.endpoint, {"file": SimpleUploadedFile("source.csv", b"id\n1\n")}, **headers,
                )
                self.assertEqual(response.status_code, 403)
                self.assertEqual(response.json()["error"]["code"], "csrf_failed")
                self.assertNotIn("attacker.example", response.content.decode())
        self.assertEqual(Dataset.objects.count(), 0)
        self.assertEqual(self.stored_files(), set())

    def test_unsupported_disguised_empty_and_binary_files_are_rejected(self):
        for filename, content in (
            ("source.exe", b"name,value\nriver,1\n"),
            ("source", b"name,value\nriver,1\n"),
            ("empty.csv", b""),
            ("blank.csv", b"\n\n"),
            ("blank.tsv", b"\t \n"),
            ("fake.xlsx", b"name,value\nriver,1\n"),
            ("fake.parquet", b"name,value\nriver,1\n"),
            ("fake.json", b"<html>not JSON</html>"),
            ("binary.csv", b"PK\x03\x04\x00binary"),
        ):
            with self.subTest(filename=filename):
                self.assert_invalid_upload(filename, content)

    def test_json_rejects_malformed_primitives_nonfinite_and_excessive_nesting(self):
        for content in (
            b'{"broken":',
            b'{"trailing":1,}',
            b'"string"',
            b"12",
            b"true",
            b"null",
            b'{"value":NaN}',
            b"[Infinity]",
            b"[-Infinity]",
            b"[" * 65 + b"0" + b"]" * 65,
            b'{"name":"\xff"}',
        ):
            with self.subTest(content=repr(content[:80])):
                self.assert_invalid_upload("invalid.json", content)

    def test_json_depth_boundary_and_delimiters_inside_strings_are_valid(self):
        for content in (
            b"[" * 64 + b"0" + b"]" * 64,
            b'{"text":"[[[[ braces { are strings } ]]]] and \\\"quotes\\\""}',
        ):
            with self.subTest(content=repr(content[:80])):
                response = self.post_file("nested.json", content)
                self.assertEqual(response.status_code, 201, response.content)

    def test_delimited_files_reject_inconsistent_rows_invalid_utf8_and_bad_quoting(self):
        for filename, content in (
            ("inconsistent.csv", b"id,value\n1,2,3\n"),
            ("inconsistent.tsv", b"id\tvalue\n1\n"),
            ("invalid.csv", b"id,value\n1,\xff\n"),
            ("invalid.tsv", b"id\tvalue\n1\t\xff\n"),
            ("unclosed.csv", b'id,value\n1,"unclosed\n'),
        ):
            with self.subTest(filename=filename):
                self.assert_invalid_upload(filename, content)

    def test_xlsx_xml_entities_and_excessive_nesting_fail_before_workbook_loading(self):
        payloads = (
            ("entity", b'<!DOCTYPE workbook [<!ENTITY marker "PRIVATE_ENTITY_VALUE">]><workbook>&marker;</workbook>'),
            ("depth", b"<node>" * 65 + b"value" + b"</node>" * 65),
        )
        for label, xml in payloads:
            with self.subTest(label=label):
                archive_buffer = io.BytesIO()
                with ZipFile(archive_buffer, "w") as archive:
                    archive.writestr("xl/workbook.xml", xml)
                with patch("datasets.validators.openpyxl.load_workbook") as load_workbook:
                    response = self.assert_invalid_upload(f"{label}.xlsx", archive_buffer.getvalue())
                load_workbook.assert_not_called()
                self.assertNotIn("PRIVATE_ENTITY_VALUE", response.content.decode())
                if label == "depth":
                    self.assertIn("XML complexity", response.json()["error"]["message"])

    def test_parquet_with_truncated_footer_is_rejected_without_storing_content(self):
        valid_parquet = self.valid_formats()["parquet"]
        self.assertTrue(valid_parquet.endswith(b"PAR1"))
        for missing_bytes in (4, 8):
            with self.subTest(missing_bytes=missing_bytes):
                self.assert_invalid_upload("truncated.parquet", valid_parquet[:-missing_bytes])

    @override_settings(DATASET_MAX_FILE_SIZE_BYTES=8)
    def test_actual_oversized_upload_chunks_are_rejected_before_service_and_storage(self):
        with patch("datasets.views.upload_dataset") as upload_service:
            response = self.post_file("large.csv", b"id,value\n" + b"1,2\n" * 20_000)
        self.assertEqual(response.status_code, 413, response.content)
        self.assertEqual(response.json()["error"]["code"], "file_too_large")
        upload_service.assert_not_called()
        self.assertEqual(Dataset.objects.count(), 0)
        self.assertEqual(self.stored_files(), set())
        listed = self.client.get(self.endpoint)
        self.assertEqual(listed.json()["limits"]["max_file_size_bytes"], 8)

    @override_settings(DATASET_MAX_FILE_SIZE_BYTES=5)
    def test_upload_at_exact_file_size_limit_is_accepted(self):
        content = b"id\n1\n"
        response = self.post_file("limit.csv", content)
        self.assertEqual(response.status_code, 201, response.content)
        self.assertEqual(response.json()["dataset"]["size_bytes"], 5)

    def test_upload_accepts_exactly_one_file_and_rejects_owner_or_other_fields(self):
        payloads = (
            {},
            {"wrong": SimpleUploadedFile("wrong.csv", b"id\n1\n")},
            {"file": [SimpleUploadedFile("one.csv", b"id\n1\n"), SimpleUploadedFile("two.csv", b"id\n2\n")]},
            {"file": SimpleUploadedFile("one.csv", b"id\n1\n"), "extra": SimpleUploadedFile("extra.csv", b"id\n2\n")},
            {"file": SimpleUploadedFile("source.csv", b"id\n1\n"), "owner": str(self.other_user.pk)},
            {"file": SimpleUploadedFile("source.csv", b"id\n1\n"), "label": "unexpected"},
        )
        for payload in payloads:
            with self.subTest(fields=list(payload)):
                response = self.client.post(
                    self.endpoint, payload, HTTP_X_CSRFTOKEN=self.token, HTTP_ORIGIN="http://testserver",
                )
                self.assertEqual(response.status_code, 400, response.content)
                self.assertEqual(response.json()["error"]["code"], "validation_error")
                self.assertIn("file", response.json()["error"]["fields"])
        self.assertEqual(Dataset.objects.count(), 0)
        self.assertEqual(self.stored_files(), set())

    def test_nonmultipart_upload_and_unsupported_methods_return_json_errors(self):
        response = self.client.post(
            self.endpoint, data='{"file":"source.csv"}', content_type="application/json",
            HTTP_X_CSRFTOKEN=self.token, HTTP_ORIGIN="http://testserver",
        )
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.json()["error"]["code"], "invalid_request")
        response = self.client.delete(
            self.endpoint, HTTP_X_CSRFTOKEN=self.token, HTTP_ORIGIN="http://testserver",
        )
        self.assertEqual(response.status_code, 405)
        self.assertEqual(response.json()["error"]["code"], "method_not_allowed")
        self.assertEqual(response["Allow"], "GET, POST")

    def test_paginated_results_follow_database_order_and_reject_invalid_pages(self):
        Dataset.objects.bulk_create([
            Dataset(owner=self.user, filename=f"source-{index}.csv", format="csv",
                    file=f"datasets/{self.user.pk}/fixture-{index}.csv", size_bytes=5, sha256="a" * 64)
            for index in range(21)
        ])
        expected = list(Dataset.objects.filter(owner=self.user).values_list("id", flat=True))
        first = self.client.get(self.endpoint, {"page": "1"})
        second = self.client.get(self.endpoint, {"page": "2"})
        self.assertEqual(first.status_code, 200)
        self.assertEqual(second.status_code, 200)
        self.assertEqual(first.json()["total_count"], 21)
        self.assertEqual(first.json()["total_pages"], 2)
        self.assertEqual(first.json()["page"], 1)
        self.assertEqual(second.json()["page"], 2)
        self.assertEqual([item["id"] for item in first.json()["datasets"]], [str(pk) for pk in expected[:20]])
        self.assertEqual([item["id"] for item in second.json()["datasets"]], [str(pk) for pk in expected[20:]])
        for page in ("", "0", "-1", "1.5", "abc", "3", "9999999999", "\u0661"):
            with self.subTest(page=page):
                response = self.client.get(self.endpoint, {"page": page})
                self.assertEqual(response.status_code, 400, response.content)
                self.assertEqual(response.json()["error"]["code"], "invalid_page")

    def test_unsafe_names_never_control_storage_path_and_uppercase_extension_is_normalized(self):
        for filename in ("../../outside.csv", r"..\..\outside.csv", "/absolute/outside.CSV"):
            with self.subTest(filename=filename):
                response = self.post_file(filename)
                self.assertEqual(response.status_code, 201, response.content)
                dataset = Dataset.objects.get(pk=response.json()["dataset"]["id"])
                self.assertNotIn("/", dataset.filename)
                self.assertNotIn("\\", dataset.filename)
                self.assertEqual(dataset.format, "csv")
                expected = f"datasets/{self.user.pk}/{dataset.pk}.csv"
                self.assertEqual(dataset.file.name, expected)
                self.assertEqual(Path(dataset.file.path).resolve(), (self.media_root / expected).resolve())
                self.assertTrue(Path(dataset.file.path).resolve().is_relative_to(self.media_root.resolve()))
                self.assertNotIn("outside", dataset.file.name)

    def test_duplicate_original_filenames_have_distinct_storage_and_preserve_both_files(self):
        contents = (b"id\n1\n", b"id\n2\n")
        identifiers = []
        for content in contents:
            response = self.post_file("same.csv", content)
            self.assertEqual(response.status_code, 201)
            identifiers.append(response.json()["dataset"]["id"])
        first, second = (Dataset.objects.get(pk=pk) for pk in identifiers)
        self.assertNotEqual(first.file.name, second.file.name)
        for dataset, content in zip((first, second), contents):
            with dataset.file.open("rb") as source:
                self.assertEqual(source.read(), content)

    def test_storage_save_failure_returns_safe_503_without_creating_metadata(self):
        self.assertEqual(self.post_file("existing.csv").status_code, 201)
        stored_before = self.stored_files()
        storage = Dataset._meta.get_field("file").storage
        with patch.object(storage, "save", side_effect=OSError("private/storage/path failed")), self.assertLogs("services.dataset_service", level="ERROR"):
            response = self.post_file("new.csv")
        self.assertEqual(response.status_code, 503)
        self.assertEqual(response.json()["error"]["code"], "upload_failed")
        self.assertNotIn("private/storage/path", response.content.decode())
        self.assertEqual(Dataset.objects.count(), 1)
        self.assertEqual(self.stored_files(), stored_before)

    def test_database_insert_failure_removes_only_the_new_file_and_transaction_remains_usable(self):
        existing = self.post_file("existing.csv")
        self.assertEqual(existing.status_code, 201)
        original = Dataset.objects.get(pk=existing.json()["dataset"]["id"])
        stored_before = self.stored_files()
        storage = Dataset._meta.get_field("file").storage
        with patch("services.dataset_service.insert_dataset", side_effect=IntegrityError("private database insert error")), patch.object(storage, "delete", wraps=storage.delete) as delete, self.assertLogs("services.dataset_service", level="ERROR"):
            response = self.post_file("new.csv")
        self.assertEqual(response.status_code, 503)
        self.assertEqual(response.json()["error"]["code"], "upload_failed")
        self.assertNotIn("private database insert error", response.content.decode())
        delete.assert_called_once()
        deleted_name = delete.call_args.args[0]
        self.assertNotEqual(deleted_name, original.file.name)
        self.assertTrue(deleted_name.startswith(f"datasets/{self.user.pk}/"))
        self.assertEqual(self.stored_files(), stored_before)
        self.assertEqual(Dataset.objects.count(), 1)
        self.assertTrue(storage.exists(original.file.name))
        response = self.post_file("after-retry.csv")
        self.assertEqual(response.status_code, 201, response.content)
        self.assertEqual(Dataset.objects.count(), 2)

    def test_database_constraints_reject_nonpositive_size_and_unsupported_format(self):
        for changes in ({"size_bytes": 0}, {"size_bytes": -1}, {"format": "xml"}):
            with self.subTest(changes=changes):
                attributes = {"owner": self.user, "filename": "invalid.csv", "format": "csv",
                              "file": "datasets/constraint-only.csv", "size_bytes": 5, "sha256": "a" * 64}
                attributes.update(changes)
                with self.assertRaises(IntegrityError):
                    with transaction.atomic():
                        Dataset.objects.create(**attributes)
                self.assertEqual(Dataset.objects.count(), 0)
